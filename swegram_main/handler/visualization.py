import json
import os
from copy import copy
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from openpyxl import worksheet

from swegram_main.data.features import Feature
from swegram_main.data.texts import Corpus
from swegram_main.handler.handler import load
from swegram_main.lib.utils import XlsxClient


class Visualization:

    def __init__(
        self, input_path: Path, language: str, output_dir: Optional[Path],
        include_tags: Optional[List[str]], exclude_tags: Optional[List[str]]
    ) -> None:
        self.language = language
        self.input_path = input_path
        include_labels = include_tags or []
        exclude_labels = exclude_tags or []
        self.labels = f"Include metadata: {' '.join(include_labels)}\n" if include_labels else ""
        self.labels += f"Exclude metadata: {' '.join(exclude_labels)}\n" if exclude_labels else ""
        self.corpus: Corpus = load(input_path, language, include_tags, exclude_tags)
        self.outdir = output_dir or Path(os.getcwd())
        os.makedirs(self.outdir, exist_ok=True)

    def filter(
        self, units: List[str], aspects: List[str],
        include_features: List[str], exclude_features: List[str],
        pprint: bool = False, save_as: str = "txt"
    ) -> None:
        self.units = units
        self.pprint = pprint
        self.aspects = aspects
        self.save_as = save_as
        data = OrderedDict()
        if "corpus" in units:
            data["corpus"] = self.filter_instance(self.corpus, include_features, exclude_features)
        if "text" in units:
            data["text"] = [
                self.filter_instance(text, include_features, exclude_features) for text in self.corpus.texts
            ]
        if "paragraph" in units:
            data["paragraph"] = [
                [
                    self.filter_instance(p, include_features, exclude_features)
                    for p in text.paragraphs
                ]
                for text in self.corpus.texts
            ]
        if "sentence" in units:
            data["sentence"] = [
                [
                    [
                        self.filter_instance(s, include_features, exclude_features)
                        for s in p.sentences
                    ] for p in text.paragraphs
                ] for text in self.corpus.texts
            ]

        self.outfile_name = self.outdir.joinpath(f"statistic-{self.input_path.with_suffix(f'.{save_as}').name}")
        if save_as == "txt" or pprint:
            self.save(data)
        if save_as == "json":
            with open(self.outfile_name, code="w", encoding="utf-8") as output_file:
                data["metadata"] = self.get_json_info()
                json_object = json.dumps(self.serialize_json_data(data), indent=4)
                output_file.write(json_object)
        elif save_as == "xlsx":
            XlsxStatisticWriter(self.outfile_name).load(self.get_json_info(), self.aspects, data)

    def serialize_json_data(self, data: Any) -> str:
        if isinstance(data, OrderedDict):
            for key, value in data.items():
                if isinstance(value, Feature):
                    data[key] = value.json
                elif isinstance(value, list):
                    data[key] = self.serialize_json_data(value)
        elif isinstance(data, list):
            for index, instance in enumerate(data):
                data[index] = self.serialize_json_data(instance)
        return data

    def append_in_text(self, content: str, mode: str = "a+") -> None:
        with open(self.outfile_name, mode, encoding="utf-8") as output_file:
            output_file.write(f"{content}\n")

    def get_info(self) -> str:
        return "Swegram statistic\n" \
               f"Time: {str(datetime.now())}\n" \
               f"Language: {self.language}\n" \
               f"Labels: {self.labels}" if self.labels else "" \
               f"Units: {self.units}\n" \
               f"Aspects: {self.aspects}\n"

    def get_json_info(self) -> Dict[str, str]:
        instance_info = {
            "Time": str(datetime.now()),
            "Language": self.language,
            "Units": self.units,
            "Aspects": self.aspects
        }
        if self.labels:
            instance_info.update({"Labels": self.labels})
        return instance_info

    def filter_instance(
        self, instance, include_features: List[str], exclude_features: List[str]
    ) -> List[Dict[str, Dict[str, Union[int, float]]]]:
        return [self.filter_aspect(instance, aspect, include_features, exclude_features) for aspect in self.aspects]

    def filter_aspect(self, instance, aspect: str, include_features: List[str], exclude_features: List[str]):
        aspect_dict = copy(getattr(instance, aspect))
        if exclude_features:
            for feature in exclude_features:
                if feature in aspect_dict:
                    del aspect_dict[feature]
        if include_features:
            for feature in aspect_dict:
                if feature not in include_features:
                    del aspect_dict[feature]
        return aspect_dict

    def save(self, data: OrderedDict) -> None:
        if self.pprint:
            print(self.get_info())
        if self.save_as == "txt":
            self.append_in_text(self.get_info(), mode="w")

        for unit in data:
            if unit == "corpus":
                self.save_instance(None, unit, data[unit])
            elif unit == "text":
                for text_index, text_instance in enumerate(data[unit], 1):
                    self.save_instance(str(text_index), unit, text_instance)
            elif unit == "paragraph":
                for ti, text_list in enumerate(data[unit], 1):
                    for pi, paragraph_instance in enumerate(text_list, 1):
                        self.save_instance(f"{ti}-{pi}", unit, paragraph_instance)
            elif unit == "sentence":
                for ti, text_list in enumerate(data[unit], 1):
                    for pi, p_list in enumerate(text_list, 1):
                        for si, sentence_instance in enumerate(p_list, 1):
                            self.save_instance(f"{ti}-{pi}-{si}", unit, sentence_instance)
            if self.pprint:
                print()

    def save_instance(self, index: Optional[str], unit: str, aspect_instances: List[OrderedDict]) -> None:
        self._save_instance_head()
        for aspect_name, instance in zip(self.aspects, aspect_instances):
            self.save_aspect_title(unit, aspect_name, index)
            for fn, f in instance.items():
                self.save_feature(fn, f)
            if self.pprint:
                print()
        if self.pprint:
            print()
        if self.save_as == "txt":
            self.append_in_text("")

    def save_aspect_title(self, unit: str, aspect_name: str, index: Optional[str]) -> None:
        title = f"{' ':>2}{'-'.join(e for e in (unit.title(), index, aspect_name) if e):>40}" \
                f"{'|':>4}{'-'*13}|{'-'*13}|{'-'*13}|"
        if self.pprint:
            print(title)
        if self.save_as == "txt":
            self.append_in_text(title)

    def save_feature(self, fn: str, f: Feature) -> None:
        def c(v: Any) -> Any:
            return v or ""

        feature = f"{' ':>2}{fn:>40}{'|':>4}{c(f.scalar):>10}{'|':>4}{c(f.mean):>10}{'|':>4}{c(f.median):>10}{'|':>4}"
        if self.pprint:
            print(feature)
        if self.save_as == "txt":
            self.append_in_text(feature)

    def _save_instance_head(self):
        af = "UNIT-ASPECT/Features"
        scalar = "Scalar"
        mean = "Mean"
        median = "Median"

        heading = f"{' ':>2}{af:>40}{'|':>4}{scalar:>10}{'|':>4}{mean:>10}{'|':>4}{median:>10}{'|':>4}"
        if self.pprint:
            print(heading)
        if self.save_as == "txt":
            self.append_in_text(heading)


class XlsxStatisticWriter(XlsxClient):

    def load(self, header: Dict[str, Union[str, List[str]]], aspects: List[str], data: Any) -> None:

        meta_sheet = self.wb["Sheet"]
        self.aspects = aspects
        meta_sheet.title = "Statistic-metadata"
        meta_sheet["A1"] = "Swegram statistics"
        for row, (key, value) in enumerate(header.items(), 2):
            self.dump_cell(meta_sheet, row, 1, key)
            if isinstance(value, str):
                self.dump_cell(meta_sheet, row, 2, value)
            elif isinstance(value, list):
                self.dump_column_list(meta_sheet, row, 2, value)

        for unit, aspect_list in data.items():
            self.load_unit(unit, aspect_list)

        self.wb.save(filename=self.output_name)

    def load_unit(self, unit: str, aspects: Any) -> None:
        sheet = self.wb.create_sheet(title=unit)
        self.dump_cell(sheet, 1, 1, f"Statistic-{unit}")
        self.load_aspects(sheet, 2, unit, aspects)

    def load_aspects(self, sheet: worksheet, row: int, unit: str, aspects: Any) -> int:
        if isinstance(aspects[0], OrderedDict):
            for aspect, features in zip(self.aspects, aspects):
                for col, content in enumerate([f"{unit}-{aspect}", "Scalar", "Mean", "Median"], 1):
                    self.dump_cell(sheet, row, col, content)
                row += 1
                for feature_name, feature in features.items():
                    self.dump_column_list(sheet, row, 1, [feature_name, feature.scalar, feature.mean, feature.median])
                    row += 1
            row += 1
        elif isinstance(aspects[0], list):
            for index, instance in enumerate(aspects, 1):
                row = self.load_aspects(sheet, row, f"{unit}-{index}", instance)
        return row
